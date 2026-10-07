"""Add more content to push report closer to ~30 pages."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")
BODY_PT = 14


def set_run_font(run, size_pt=BODY_PT, bold=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        run = new_para.add_run(text)
        set_run_font(run, BODY_PT)
    return new_para


def add_body_after(paragraph: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(paragraph, text, style="Normal")
    try:
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.5
    except Exception:
        pass
    return p


def add_heading_after(paragraph: Paragraph, text: str) -> Paragraph:
    p = insert_paragraph_after(paragraph, "", style="Heading 2")
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    set_run_font(run, BODY_PT, bold=True)
    return p


def find_contains(doc, needle):
    for i, p in enumerate(doc.paragraphs):
        if needle in (p.text or ""):
            return i
    return None


def find_startswith(doc, needle):
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip().startswith(needle):
            return i
    return None


EXTRA = [
    ("H", "4.35 Extended Technical Discussion – Streaming Parameters"),
    ("N", "The MJPEG stream from ESP32-CAM provides a continuous sequence of JPEG frames. For detection, the browser does not need every possible camera frame rate; a moderate update rate is enough if detection latency remains acceptable for the use case. During testing, the stream was kept stable enough for COCO-SSD to produce usable boxes when a dog entered the view. Resolution selection is a trade-off: higher resolution improves detail but increases bandwidth and processing cost, while lower resolution is faster but may miss distant subjects."),
    ("N", "In practical terms, the camera should be mounted so that the car roof or target zone occupies a meaningful portion of the frame. This improves the apparent size of the dog and increases the chance that COCO-SSD will assign a confidence score above the chosen threshold. Mounting height and angle were treated as part of system tuning, not only as mechanical installation."),
    ("N", "Wi-Fi quality affects both stream smoothness and control API responsiveness. If the signal is weak, frames may freeze briefly and detection can lag. For demonstration and evaluation, the board and viewing laptop were kept on a reliable local network. This is an important deployment lesson from the internship project."),
    ("H", "4.36 Extended Technical Discussion – Confidence Threshold Selection"),
    ("N", "A confidence threshold of approximately 55 percent was selected after practical trials. A very high threshold reduces false positives but increases missed detections. A very low threshold increases alerts, including less reliable ones. For a deterrent/alert prototype, missing some uncertain detections is often preferable to frequent false relay triggers, yet the threshold should not be so high that obvious dogs are ignored."),
    ("N", "Threshold selection is therefore an applied Machine Learning decision made at system level. The model itself was not retrained for this final path; instead, decision policy around the model was tuned. This is a common industrial pattern: reuse a capable pre-trained model and specialize behavior through filters, thresholds, zones, and timing logic."),
    ("N", "The displayed percentage on the green box helps the operator understand why an alert occurred. If scores are repeatedly near the threshold, the operator can improve lighting, camera distance, or zone placement before considering model replacement."),
    ("H", "4.37 Extended Technical Discussion – ROI Geometry Logic"),
    ("N", "The DETECT ZONE is a rectangle defined by user drag gestures on the camera canvas. Internally, the zone can be stored as minimum and maximum x/y coordinates. For each accepted dog box, the center point is estimated as the midpoint of the box. If that point lies within the rectangle, the detection is valid for alerting; otherwise it is ignored for actuation even if a box is drawn for visualization policy choices."),
    ("N", "Using the center point is simple and robust for this prototype. Alternatives include requiring the full box to be inside the zone or requiring a minimum overlap percentage (Intersection-over-Union style checks against the zone). Center-point gating was sufficient for the training goal and easy to explain during demonstration."),
    ("N", "Clearing the zone restores full-frame monitoring. This is useful when the operator wants temporary wide-area observation without redesigning camera placement. The feature therefore increases operational flexibility without firmware changes."),
    ("H", "4.38 Relay Hold State Machine"),
    ("N", "The alert behavior can be described as a small state machine with states such as IDLE, ALERT_HOLD, and RECHECK. In IDLE, the system waits for a valid dog event. On event, it enters ALERT_HOLD, sets relay ON, and starts a 7-second timer. During ALERT_HOLD, temporary detection gaps do not immediately clear the output. When the timer ends, the system enters RECHECK: if a valid dog is still present, it returns to ALERT_HOLD; otherwise it turns outputs OFF and returns to IDLE."),
    ("N", "Expressing control logic as states reduces ad-hoc bugs and makes timing behavior predictable. It also helps in viva explanation: one can describe exactly when the relay should be ON and why brief model flicker does not create relay chatter."),
    ("N", "This state-machine thinking connects embedded control concepts with AI inference outputs and is an important engineering takeaway from the project."),
    ("H", "4.39 Power, Boot and Reliability Notes"),
    ("N", "ESP32-CAM modules are sensitive to power quality, especially when the camera and Wi-Fi are active. An inadequate supply can cause brownouts, reboots, or camera initialization failures. During development, a stable 5V supply and correct upload wiring were essential. After boot, waiting for Wi-Fi association and confirming the IP on Serial Monitor remained the standard bring-up checklist."),
    ("N", "Firmware reliability also depends on not blocking the networking stack for long periods and on using appropriate APIs for LED control. The project intentionally kept heavy inference off the microcontroller to protect streaming reliability."),
    ("N", "From a maintenance view, separating UI/detection (index_dog.h) from HTTP/GPIO (app_httpd.cpp) and setup (CameraWebServer.ino) made iterative fixes easier. When a bug appeared, it was usually clear whether the fault was in detection JavaScript, HTTP handling, or pin wiring."),
    ("H", "4.40 Data Privacy and Deployment Checklist"),
    ("N", "Before any longer outdoor deployment, a checklist should include: camera facing only the intended parking area, avoiding unnecessary capture of public pathways if possible, securing the Wi-Fi credentials in firmware, placing the board in a weather-aware enclosure, verifying relay load ratings, and confirming that deterrent actions are proportionate and lawful."),
    ("N", "The internship version is an academic prototype. Moving from prototype to product would additionally require hardening, logging, remote update strategy, and more rigorous accuracy evaluation across days and lighting conditions."),
    ("N", "Even so, the current system already demonstrates the core value proposition: automatic vision-based awareness plus timely actuation signaling."),
    ("H", "4.41 Work Completed – Consolidated Bullet Summary"),
    ("N", "For quick reference, the consolidated list of completed project work is:"),
    ("B", "Studied ML/DL fundamentals during industrial training sessions."),
    ("B", "Selected a real problem: stray dog presence near parked cars."),
    ("B", "Implemented ESP32-CAM live streaming web monitor."),
    ("B", "Integrated TensorFlow.js COCO-SSD object detection."),
    ("B", "Restricted alerts to dog class with confidence threshold."),
    ("B", "Added ROI detect-zone selection and clear-zone control."),
    ("B", "Implemented relay control on GPIO12 with 7-second hold/recheck."),
    ("B", "Implemented flash LED control on GPIO4 (auto/manual)."),
    ("B", "Performed functional tests for in-zone/out-zone and non-dog cases."),
    ("B", "Documented architecture, figures, and results in this report."),
    ("H", "4.42 Sample Demonstration Script Used During Testing"),
    ("N", "A typical demonstration sequence used during mentor/self testing was as follows. First, power the ESP32-CAM and confirm Serial Monitor connection message and IP address. Second, open the IP in Chrome/Edge and confirm live video. Third, wait for model load indication/ready state. Fourth, draw a DETECT ZONE over the critical car region. Fifth, introduce a dog target into the zone and observe green box, confidence text, relay ON and flash behavior. Sixth, move the target outside the zone and confirm alert does not remain justified after recheck. Seventh, clear the zone and show full-frame mode."),
    ("N", "This script ensures every major feature is exercised in a short session and is useful for viva/practical evaluation."),
    ("H", "4.43 Role of Each Major File in the Codebase"),
    ("N", "CameraWebServer.ino is the entry firmware file. It configures the camera pins for the AI-Thinker board variant, connects to Wi-Fi, initializes relay direction, sets up flash LED hardware access, and starts the web server. Without successful execution of this file’s setup path, neither streaming nor APIs are available."),
    ("N", "app_httpd.cpp implements the HTTP side: serving the HTML UI, providing the MJPEG stream endpoint, and handling /relay and /led requests from the browser. This file is the bridge between web commands and GPIO actions."),
    ("N", "index_dog.h stores the front-end application: layout, styling, stream element, TensorFlow.js usage, dog filtering, ROI interactions, and alert timing. board_config.h keeps hardware constants centralized. Optional Python scripts under learning/experiment folders supported ML practice but are not required for the final live browser detection path."),
    ("H", "4.44 End of Extended Project Discussion"),
    ("N", "The extended sections above provide additional engineering depth around streaming, thresholds, ROI geometry, alert state behavior, power/reliability, privacy/deployment, consolidated work summary, demonstration flow, and code file roles. They complement the earlier implementation narrative and figures so that the project chapter is complete for academic submission."),
]


def main():
    doc = Document(str(SRC))
    if find_contains(doc, "4.35 Extended Technical Discussion") is not None:
        print("Already expanded once")
    else:
        # insert before CHAPTER 5
        ch5 = find_startswith(doc, "CHAPTER 5:")
        if ch5 is None:
            raise SystemExit("CHAPTER 5 not found")
        j = ch5 - 1
        while j > 0 and not (doc.paragraphs[j].text or "").strip():
            j -= 1
        cur = doc.paragraphs[j]
        for kind, text in EXTRA:
            if kind == "H":
                cur = add_heading_after(cur, text)
            elif kind == "B":
                cur = add_body_after(cur, "• " + text)
            else:
                cur = add_body_after(cur, text)

    # Mild spacing bump on Normal paragraphs without touching fonts/colors
    # (already set on new paras)

    doc.save(str(SRC))
    words = sum(len((p.text or "").split()) for p in doc.paragraphs)
    print("words", words)
    print("approx_pages_300w", round(words / 300, 1))
    print("approx_pages_250w", round(words / 250, 1))


if __name__ == "__main__":
    main()
