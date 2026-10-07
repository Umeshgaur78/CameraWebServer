"""Apply real H1 page numbers; nudge to ~30 pages; verify LoF not blank."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")
BODY_PT = 14

# From Word Heading 1 walk (print layout)
PAGES = {
    "CANDIDATE'S DECLARATION": 2,
    "CERTIFICATE": 3,
    "ACKNOWLEDGEMENT": 4,
    "ABSTRACT": 5,
    "CONTENTS": 8,
    "List of Figures": 9,
    "List of Tables": 10,
    "CHAPTER 1: INTRODUCTION": 11,
    "CHAPTER 2: COMPANY INFRASTRUCTURE": 13,
    "CHAPTER 3: TRAINING ATTENDED": 15,
    "CHAPTER 4: SYSTEMS/PROJECT DEVELOPMENT": 17,
    "CHAPTER 5: RESULTS AND CONCLUSION": 27,
    "REFERENCES": 29,
    # figures fall in ch4 before ch5
    "Fig. 4.1": 22,
    "Fig. 4.2": 23,
    "Fig. 4.3": 24,
    "Fig. 4.4": 25,
    "Table 4.1": 18,
    "Table 4.2": 21,
}


def set_run_font(run, size_pt=BODY_PT, bold=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
        set_run_font(paragraph.runs[0], BODY_PT)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, BODY_PT)


def dotted(title, page, width=88):
    page = str(page)
    dots = max(3, width - len(title) - len(page))
    return f"{title} {'.' * dots} {page}"


def find_exact(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip() == text:
            return i
    return None


def find_startswith(doc, needle):
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip().startswith(needle):
            return i
    return None


def main():
    doc = Document(str(SRC))

    # Update TOC lines only between CONTENTS and List of Figures heading
    contents = find_exact(doc, "CONTENTS")
    lof = find_exact(doc, "List of Figures")
    toc_lines = [
        ("Candidate's Declaration", PAGES["CANDIDATE'S DECLARATION"]),
        ("Certificate", PAGES["CERTIFICATE"]),
        ("Acknowledgement", PAGES["ACKNOWLEDGEMENT"]),
        ("Abstract", PAGES["ABSTRACT"]),
        ("Contents", PAGES["CONTENTS"]),
        ("List of Figures", PAGES["List of Figures"]),
        ("List of Tables", PAGES["List of Tables"]),
        ("Chapter 1: Introduction", PAGES["CHAPTER 1: INTRODUCTION"]),
        ("Chapter 2: Company Infrastructure", PAGES["CHAPTER 2: COMPANY INFRASTRUCTURE"]),
        ("Chapter 3: Training Attended", PAGES["CHAPTER 3: TRAINING ATTENDED"]),
        ("Chapter 4: Systems / Project Development", PAGES["CHAPTER 4: SYSTEMS/PROJECT DEVELOPMENT"]),
        ("Chapter 5: Results and Conclusion", PAGES["CHAPTER 5: RESULTS AND CONCLUSION"]),
        ("References", PAGES["REFERENCES"]),
    ]
    if contents is not None and lof is not None:
        for i in range(contents + 1, lof):
            t = (doc.paragraphs[i].text or "").strip()
            for title, page in toc_lines:
                if t.startswith(title):
                    replace_paragraph_text(doc.paragraphs[i], dotted(title, page))
                    break

    # LoF / LoT tables
    if len(doc.tables) >= 1:
        t0 = doc.tables[0]
        data = [
            ("Fig. 4.1", "ESP32-CAM hardware concept", str(PAGES["Fig. 4.1"])),
            ("Fig. 4.2", "System architecture of dog detection alert system", str(PAGES["Fig. 4.2"])),
            ("Fig. 4.3", "Detection and relay-hold workflow", str(PAGES["Fig. 4.3"])),
            ("Fig. 4.4", "Monitoring UI with DETECT ZONE and dog box", str(PAGES["Fig. 4.4"])),
        ]
        while len(t0.rows) < 5:
            t0.add_row()
        for i, (a, b, c) in enumerate(data, start=1):
            t0.rows[i].cells[0].text = a
            t0.rows[i].cells[1].text = b
            t0.rows[i].cells[2].text = c

    if len(doc.tables) >= 2:
        t1 = doc.tables[1]
        while len(t1.rows) < 3:
            t1.add_row()
        t1.rows[0].cells[0].text = "Table No."
        t1.rows[1].cells[0].text = "Table 4.1"
        t1.rows[1].cells[1].text = "Tools and technologies used in the project"
        t1.rows[1].cells[2].text = str(PAGES["Table 4.1"])
        t1.rows[2].cells[0].text = "Table 4.2"
        t1.rows[2].cells[1].text = "GPIO pin mapping (Relay/Flash)"
        t1.rows[2].cells[2].text = str(PAGES["Table 4.2"])

    # Add a little content to reach ~30 pages (currently 29)
    if find_startswith(doc, "4.19 Closing Note") is None:
        ch5 = find_startswith(doc, "CHAPTER 5:")
        j = ch5 - 1
        while j > 0 and not (doc.paragraphs[j].text or "").strip():
            j -= 1
        cur = doc.paragraphs[j]

        def add_after(paragraph, text, style="Normal", bold=False):
            new_p = OxmlElement("w:p")
            paragraph._p.addnext(new_p)
            para = Paragraph(new_p, paragraph._parent)
            try:
                para.style = doc.styles[style]
            except Exception:
                pass
            run = para.add_run(text)
            set_run_font(run, BODY_PT, bold=bold)
            try:
                para.paragraph_format.line_spacing = 1.5
                para.paragraph_format.space_after = Pt(8)
            except Exception:
                pass
            return para

        cur = add_after(cur, "4.19 Closing Note of Project Chapter", style="Heading 2", bold=True)
        cur = add_after(
            cur,
            "This chapter presented the complete project work done during industrial training: problem understanding, "
            "system design, ESP32-CAM implementation, COCO-SSD based dog detection, ROI based monitoring, relay and flash "
            "control, figures, and testing notes. The following chapter summarizes results and conclusion.",
        )

    doc.save(str(SRC))
    print("Updated TOC/LoF/LoT page numbers and saved.")

    # quick verify front sequence
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if i > 95:
            break
        if t in ("CONTENTS", "List of Figures", "List of Tables") or t.startswith(
            ("Chapter 1", "Chapter 4", "CHAPTER 1", "Candidate", "ABSTRACT")
        ):
            print(i, t[:90])


if __name__ == "__main__":
    main()
