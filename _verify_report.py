from docx import Document

p = r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx"
d = Document(p)
words = 0
imgs = 0
for para in d.paragraphs:
    t = para.text or ""
    words += len(t.split())
    if para._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'):
        imgs += 1
print("paras", len(d.paragraphs))
print("words", words)
print("approx_pages_350w", round(words / 350, 1))
print("images", imgs)
for i, para in enumerate(d.paragraphs):
    if (para.text or "").strip() == "CERTIFICATE":
        print("cert_at", i)
        for j in range(i, min(i + 6, len(d.paragraphs))):
            t = (d.paragraphs[j].text or "").strip()
            has = bool(d.paragraphs[j]._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
            print(" ", j, repr(t[:70]), "IMG" if has else "")
        break
print("has_4.4", any("4.4 Tools" in (x.text or "") for x in d.paragraphs))
print("has_fig41", any("Fig. 4.1" in (x.text or "") for x in d.paragraphs))
print("has_abstract_done", any("successfully developed" in (x.text or "") for x in d.paragraphs))
