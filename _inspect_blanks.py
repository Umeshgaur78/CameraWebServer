from docx import Document
from docx.oxml.ns import qn

doc = Document(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")

print("Between ABSTRACT and CONTENTS:")
in_range = False
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    st = p.style.name if p.style else ""
    pb = False
    if p._p.pPr is not None:
        for c in p._p.pPr:
            if c.tag == qn("w:pageBreakBefore"):
                pb = True
    has_img = bool(p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
    if t == "ABSTRACT":
        in_range = True
    if in_range:
        print(f"{i:3d} PB={pb} img={has_img} [{st}] {repr(t[:80])}")
    if t == "CONTENTS":
        break

print("\nBetween LoF and CHAPTER1:")
in_range = False
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    st = p.style.name if p.style else ""
    pb = False
    if p._p.pPr is not None:
        for c in p._p.pPr:
            if c.tag == qn("w:pageBreakBefore"):
                pb = True
    has_img = bool(p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
    if t == "List of Figures":
        in_range = True
    if in_range:
        print(f"{i:3d} PB={pb} img={has_img} [{st}] {repr(t[:80])}")
    if t.startswith("CHAPTER 1"):
        break

# body children around this area
print("\nBody XML around front:")
body = doc.element.body
para_map = {p._p: ((p.text or "").strip()[:40], i) for i, p in enumerate(doc.paragraphs)}
for child in body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "p":
        t, i = para_map.get(child, ("?", -1))
        if i >= 56 and i <= 85:
            print(f"P[{i}] {t}")
    elif tag == "tbl":
        print("TABLE")
