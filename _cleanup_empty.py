from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx"
doc = Document(SRC)

# Remove empty Heading paragraphs that create blank/odd pages
removed = 0
for p in list(doc.paragraphs):
    t = (p.text or "").strip()
    st = p.style.name if p.style else ""
    has_img = bool(p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
    if st.startswith("Heading") and not t and not has_img:
        el = p._p
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
            removed += 1

# Ensure LoF/LoT have a one-line helper under heading if only blank before next heading
# (table already present in XML after LoF)

# Add final short para before ch5 if needed for ~30 pages
marker = "This completes the project development documentation for the internship."
if not any(marker in (p.text or "") for p in doc.paragraphs):
    ch5 = None
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").startswith("CHAPTER 5:"):
            ch5 = i
            break
    j = ch5 - 1
    while j > 0 and not (doc.paragraphs[j].text or "").strip():
        j -= 1
    cur = doc.paragraphs[j]
    new_p = OxmlElement("w:p")
    cur._p.addnext(new_p)
    para = Paragraph(new_p, cur._parent)
    try:
        para.style = doc.styles["Normal"]
    except Exception:
        pass
    run = para.add_run(marker + " The next chapter presents module-wise results, lessons learned, and conclusion.")
    run.font.size = Pt(14)
    try:
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.space_after = Pt(12)
    except Exception:
        pass

# Sync LoF pages once more
if len(doc.tables) >= 1:
    t0 = doc.tables[0]
    data = [
        ("Fig. 4.1", "ESP32-CAM hardware concept", "23"),
        ("Fig. 4.2", "System architecture of dog detection alert system", "23"),
        ("Fig. 4.3", "Detection and relay-hold workflow", "24"),
        ("Fig. 4.4", "Monitoring UI with DETECT ZONE and dog box", "24"),
    ]
    while len(t0.rows) < 5:
        t0.add_row()
    for i, (a, b, c) in enumerate(data, start=1):
        t0.rows[i].cells[0].text = a
        t0.rows[i].cells[1].text = b
        t0.rows[i].cells[2].text = c

if len(doc.tables) >= 2:
    t1 = doc.tables[1]
    while len(t1.rows) < 3:
        t1.add_row()
    t1.rows[0].cells[0].text = "Table No."
    t1.rows[1].cells[0].text = "Table 4.1"
    t1.rows[1].cells[1].text = "Tools and technologies used in the project"
    t1.rows[1].cells[2].text = "18"
    t1.rows[2].cells[0].text = "Table 4.2"
    t1.rows[2].cells[1].text = "GPIO pin mapping (Relay/Flash)"
    t1.rows[2].cells[2].text = "21"

doc.save(SRC)
print("removed_empty_headings", removed)
print("saved")
