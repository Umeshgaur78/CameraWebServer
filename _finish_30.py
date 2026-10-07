from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph
import subprocess

SRC = r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx"
doc = Document(SRC)


def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
        paragraph.runs[0].font.size = Pt(14)
    else:
        run = paragraph.add_run(text)
        run.font.size = Pt(14)


def dotted(title, page, width=88):
    page = str(page)
    dots = max(3, width - len(title) - len(page))
    return f"{title} {'.' * dots} {page}"


def add_after(paragraph, text, style="Normal", bold=False):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    try:
        para.style = doc.styles[style]
    except Exception:
        pass
    run = para.add_run(text)
    run.font.size = Pt(14)
    run.font.bold = bold
    try:
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.space_after = Pt(10)
    except Exception:
        pass
    return para


# Add one page-equivalent content before REFERENCES if not present
if not any("5.6 Overall Internship Outcome" in (p.text or "") for p in doc.paragraphs):
    ref_i = None
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip() == "REFERENCES":
            ref_i = i
            break
    j = ref_i - 1
    while j > 0 and not (doc.paragraphs[j].text or "").strip():
        j -= 1
    cur = doc.paragraphs[j]
    cur = add_after(cur, "5.6 Overall Internship Outcome", style="Heading 2", bold=True)
    cur = add_after(
        cur,
        "The industrial training provided structured exposure to Artificial Intelligence and Machine Learning fundamentals "
        "and a practical IoT vision project. By the end of the training, a complete ESP32-CAM based dog detection and alert "
        "system was implemented and demonstrated.",
    )
    cur = add_after(
        cur,
        "Key completed outcomes include live camera streaming, browser-based COCO-SSD object detection, dog-only filtering, "
        "user-selectable detect zone, relay control with 7-second hold, and flash LED control. These outcomes together satisfy "
        "the applied learning objective of the internship.",
    )
    cur = add_after(
        cur,
        "The report therefore records both the organizational training context and the concrete project contribution. "
        "Future enhancements can improve model specialization and outdoor robustness, while the present system remains a "
        "complete and evaluable prototype for academic submission.",
    )

doc.save(SRC)

# Refresh Word page map via headings script output
r = subprocess.run(
    [
        "powershell",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        r"c:\Users\umesh\OneDrive\Documents\Arduino\CameraWebServer\_word_pages3.ps1",
    ],
    capture_output=True,
    text=True,
)
print(r.stdout)

pages = {}
for line in r.stdout.splitlines():
    if line.startswith("H1_p"):
        # H1_p11:CHAPTER 1: INTRODUCTION
        try:
            left, title = line.split(":", 1)
            pg = int(left.replace("H1_p", ""))
            pages[title.strip()] = pg
        except Exception:
            pass
    if line.startswith("TOTAL_PAGES="):
        pages["TOTAL"] = int(line.split("=", 1)[1])

print("Parsed H1 pages:", pages)

# Remap TOC using parsed pages
doc = Document(SRC)
mapping = {
    "Candidate's Declaration": pages.get("CANDIDATE'S DECLARATION", 2),
    "Certificate": pages.get("CERTIFICATE", 3),
    "Acknowledgement": pages.get("ACKNOWLEDGEMENT", 4),
    "Abstract": pages.get("ABSTRACT", 5),
    "Contents": pages.get("CONTENTS", 8),
    "List of Figures": pages.get("List of Figures", 9),
    "List of Tables": pages.get("List of Tables", 10),
    "Chapter 1: Introduction": pages.get("CHAPTER 1: INTRODUCTION", 11),
    "Chapter 2: Company Infrastructure": pages.get("CHAPTER 2: COMPANY INFRASTRUCTURE", 13),
    "Chapter 3: Training Attended": pages.get("CHAPTER 3: TRAINING ATTENDED", 15),
    "Chapter 4: Systems / Project Development": pages.get("CHAPTER 4: SYSTEMS/PROJECT DEVELOPMENT", 17),
    "Chapter 5: Results and Conclusion": pages.get("CHAPTER 5: RESULTS AND CONCLUSION", 27),
    "References": pages.get("REFERENCES", 30),
}

contents = lof = None
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    if t == "CONTENTS":
        contents = i
    if t == "List of Figures" and contents is not None and i > contents:
        lof = i
        break

if contents is not None and lof is not None:
    for i in range(contents + 1, lof):
        t = (doc.paragraphs[i].text or "").strip()
        for title, page in mapping.items():
            if t.startswith(title):
                replace_paragraph_text(doc.paragraphs[i], dotted(title, page))
                break

doc.save(SRC)
print("TOTAL", pages.get("TOTAL"))
print("Done")
