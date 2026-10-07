from docx import Document
from docx.oxml.ns import qn

doc = Document(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")
print("Front matter sequence:")
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    if not t:
        continue
    if i > 120:
        break
    st = p.style.name if p.style else ""
    if st.startswith("Heading") or t.startswith(("Chapter", "List of", "CONTENTS", "ABSTRACT", "CERTIFICATE", "ACKNOWLEDGEMENT", "CANDIDATE", "CHAPTER", "Fig.", "Table 4.", "4.1", "4.17", "4.18", "5.")):
        pb = False
        if p._p.pPr is not None:
            for c in p._p.pPr:
                if c.tag == qn("w:pageBreakBefore"):
                    pb = True
        print(f"{i:3d} {'PB' if pb else '  '} [{st}] {t[:90]}")

print("\nBody tables:")
body = doc.element.body
para_map = {p._p: i for i, p in enumerate(doc.paragraphs)}
for child in body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "tbl":
        print("TABLE near para", end=" ")
        # previous sibling
        prev = child.getprevious()
        while prev is not None and prev.tag.split("}")[-1] != "p":
            prev = prev.getprevious()
        if prev is not None:
            print(para_map.get(prev), (prev.xpath("string(.)") or "")[:50])
        else:
            print("?")

print("words", sum(len((p.text or "").split()) for p in doc.paragraphs))
# chapter 4 sections present
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    if t.startswith("4.") and (p.style and p.style.name.startswith("Heading")):
        print("SEC", t[:70])
