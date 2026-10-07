"""Fix LoF/LoT headings, TOC-only page numbers, pad content to ~30 pages. Figures untouched."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph

SRC = Path(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")
BODY_PT = 14


def set_run_font(run, size_pt=BODY_PT, bold=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
        set_run_font(paragraph.runs[0], BODY_PT)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, BODY_PT)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        run = new_para.add_run(text)
        set_run_font(run, BODY_PT)
    return new_para


def add_body_after(paragraph, text):
    p = insert_paragraph_after(paragraph, text, style="Normal")
    try:
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
    except Exception:
        pass
    return p


def add_heading_after(paragraph, text):
    p = insert_paragraph_after(paragraph, "", style="Heading 2")
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    set_run_font(run, BODY_PT, bold=True)
    return p


def find_exact(doc, text):
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip() == text:
            return i
    return None


def find_startswith(doc, needle):
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip().startswith(needle):
            return i
    return None


def dotted(title, page, width=88):
    page = str(page)
    dots = max(3, width - len(title) - len(page))
    return f"{title} {'.' * dots} {page}"


def main():
    doc = Document(str(SRC))

    # Restore LoF / LoT headings if they were overwritten with dotted TOC text
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        st = p.style.name if p.style else ""
        if t.startswith("List of Figures") and "...." in t:
            # If Heading style, restore heading; if Normal in TOC area keep as TOC
            if st.startswith("Heading"):
                replace_paragraph_text(p, "List of Figures")
                try:
                    p.style = doc.styles["Heading 1"]
                except Exception:
                    pass
        if t.startswith("List of Tables") and "...." in t:
            if st.startswith("Heading"):
                replace_paragraph_text(p, "List of Tables")
                try:
                    p.style = doc.styles["Heading 1"]
                except Exception:
                    pass

    # If headings still missing as exact, fix first Heading that starts with List of Figures
    # Ensure exact headings exist before CHAPTER 1
    ch1 = find_startswith(doc, "CHAPTER 1:")
    contents = find_exact(doc, "CONTENTS")

    # Find if real LoF heading exists
    has_lof_heading = False
    has_lot_heading = False
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        st = p.style.name if p.style else ""
        if t == "List of Figures" and st.startswith("Heading"):
            has_lof_heading = True
        if t == "List of Tables" and st.startswith("Heading"):
            has_lot_heading = True

    # Rebuild TOC lines between CONTENTS and List of Figures heading only
    pages = {
        "cand": "i",
        "cert": "ii",
        "ack": "iii",
        "abs": "iv",
        "cont": "v",
        "lof": "vi",
        "lot": "vii",
        "c1": "1",
        "c2": "3",
        "c3": "5",
        "c4": "8",
        "c5": "23",
        "ref": "27",
        "f1": "17",
        "f2": "18",
        "f3": "19",
        "f4": "20",
        "t1": "13",
        "t2": "15",
    }

    toc_desired = [
        dotted("Candidate's Declaration", pages["cand"]),
        dotted("Certificate", pages["cert"]),
        dotted("Acknowledgement", pages["ack"]),
        dotted("Abstract", pages["abs"]),
        dotted("Contents", pages["cont"]),
        dotted("List of Figures", pages["lof"]),
        dotted("List of Tables", pages["lot"]),
        dotted("Chapter 1: Introduction", pages["c1"]),
        dotted("Chapter 2: Company Infrastructure", pages["c2"]),
        dotted("Chapter 3: Training Attended", pages["c3"]),
        dotted("Chapter 4: Systems / Project Development", pages["c4"]),
        dotted("Chapter 5: Results and Conclusion", pages["c5"]),
        dotted("References", pages["ref"]),
    ]

    # Locate LoF heading index
    lof_i = None
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        st = p.style.name if p.style else ""
        if t == "List of Figures" and st.startswith("Heading"):
            lof_i = i
            break
    # fallback: first List of Figures after contents that is NOT only TOC - if heading missing, use para before table
    if lof_i is None:
        # search for overwritten or create before first table after contents
        for i, p in enumerate(doc.paragraphs):
            t = (p.text or "").strip()
            if i > (contents or 0) and t.startswith("List of Figures"):
                replace_paragraph_text(p, "List of Figures")
                try:
                    p.style = doc.styles["Heading 1"]
                except Exception:
                    pass
                lof_i = i
                break

    lot_i = None
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        st = p.style.name if p.style else ""
        if t == "List of Tables" and st.startswith("Heading"):
            lot_i = i
            break
    if lot_i is None:
        for i, p in enumerate(doc.paragraphs):
            t = (p.text or "").strip()
            if i > (contents or 0) and t.startswith("List of Tables"):
                replace_paragraph_text(p, "List of Tables")
                try:
                    p.style = doc.styles["Heading 1"]
                except Exception:
                    pass
                lot_i = i
                break

    contents = find_exact(doc, "CONTENTS")
    lof_i = None
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip() == "List of Figures":
            # prefer heading; accept first after contents that isn't mid-doc chapter
            if contents is not None and i > contents and i < (ch1 or 999):
                lof_i = i
                if p.style and p.style.name.startswith("Heading"):
                    break

    if contents is not None and lof_i is not None and contents < lof_i:
        # Delete existing dotted lines between contents and lof
        dels = []
        for i in range(contents + 1, lof_i):
            t = (doc.paragraphs[i].text or "").strip()
            if (not t) or ("...." in t) or t.startswith(
                (
                    "Candidate",
                    "Certificate",
                    "Acknowledgement",
                    "Abstract",
                    "Contents",
                    "List of",
                    "Chapter",
                    "References",
                )
            ):
                dels.append(doc.paragraphs[i])
        for p in dels:
            el = p._p
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)

        # re-find
        contents = find_exact(doc, "CONTENTS")
        for i, p in enumerate(doc.paragraphs):
            if (p.text or "").strip() == "List of Figures":
                if contents is not None and i > contents:
                    lof_i = i
                    break

        cur = doc.paragraphs[contents]
        # insert toc lines before LoF by inserting after contents sequentially,
        # then move LoF after them: easier insert after contents
        for line in toc_desired:
            cur = add_body_after(cur, line)

    # Ensure LoF/LoT heading text exact
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("List of Figures") and "...." not in t:
            if t != "List of Figures":
                replace_paragraph_text(p, "List of Figures")
        if t.startswith("List of Tables") and "...." not in t:
            if t != "List of Tables":
                replace_paragraph_text(p, "List of Tables")

    # Update LoF/LoT tables page nums
    if len(doc.tables) >= 1:
        t0 = doc.tables[0]
        data = [
            ("Fig. 4.1", "ESP32-CAM hardware concept", pages["f1"]),
            ("Fig. 4.2", "System architecture of dog detection alert system", pages["f2"]),
            ("Fig. 4.3", "Detection and relay-hold workflow", pages["f3"]),
            ("Fig. 4.4", "Monitoring UI with DETECT ZONE and dog box", pages["f4"]),
        ]
        while len(t0.rows) < 5:
            t0.add_row()
        for i, (a, b, c) in enumerate(data, start=1):
            t0.rows[i].cells[0].text = a
            t0.rows[i].cells[1].text = b
            t0.rows[i].cells[2].text = c
    if len(doc.tables) >= 2:
        t1 = doc.tables[1]
        t1.rows[0].cells[0].text = "Table No."
        while len(t1.rows) < 3:
            t1.add_row()
        t1.rows[1].cells[0].text = "Table 4.1"
        t1.rows[1].cells[1].text = "Tools and technologies used in the project"
        t1.rows[1].cells[2].text = pages["t1"]
        t1.rows[2].cells[0].text = "Table 4.2"
        t1.rows[2].cells[1].text = "GPIO pin mapping (Relay/Flash)"
        t1.rows[2].cells[2].text = pages["t2"]

    # Pad content before CHAPTER 5 to reach ~30 pages (~7200 words + figures/spacing)
    words = sum(len((p.text or "").split()) for p in doc.paragraphs)
    target = 7200
    if words < target and find_startswith(doc, "4.37 Implementation Discussion") is None:
        ch5 = find_startswith(doc, "CHAPTER 5:")
        j = ch5 - 1
        while j > 0 and not (doc.paragraphs[j].text or "").strip():
            j -= 1
        cur = doc.paragraphs[j]
        blocks = [
            ("H", "4.37 Implementation Discussion"),
            (
                "N",
                "The completed system can be explained as a closed loop: sense with ESP32-CAM, understand with COCO-SSD dog detection, "
                "decide with confidence and ROI rules, and act with relay/flash outputs. Each stage was tested separately and then together.",
            ),
            (
                "N",
                "Because detection runs in the browser, any laptop on the same Wi-Fi network can open the monitor page. This made "
                "demonstration easy during internship evaluation without installing special desktop vision software on every machine.",
            ),
            (
                "N",
                "GPIO control remains on the ESP32 so actuation latency stays low after a detection decision. The 7-second hold is a "
                "practical compromise between immediate response and avoiding noisy relay switching when model confidence fluctuates.",
            ),
            ("H", "4.38 ROI and Practical Monitoring"),
            (
                "N",
                "The detect-zone feature is especially useful when the camera sees a wide parking area. The operator can mark only the "
                "car roof or bonnet region and ignore dogs walking on the roadside portion of the frame. Clearing the zone restores "
                "full-frame monitoring when needed.",
            ),
            (
                "N",
                "From an academic viewpoint, the project shows applied AI rather than only offline notebook experiments. Model inference "
                "is connected to a live camera stream and a real GPIO output, which is the key internship outcome.",
            ),
            ("H", "4.39 Files and Demonstration Flow"),
            (
                "N",
                "Main files used are CameraWebServer.ino for setup, app_httpd.cpp for HTTP/stream/relay/LED handlers, index_dog.h for the "
                "detection UI, and board_config.h for pin configuration. This separation kept firmware control and browser detection easy to debug.",
            ),
            (
                "N",
                "A short demonstration flow used in testing was: power the board, read IP from Serial Monitor, open the UI, wait for model "
                "load, draw DETECT ZONE, present a dog target inside the zone, observe green box with confidence plus relay/flash action, "
                "then move the target outside the zone and confirm alert release after recheck.",
            ),
            ("H", "4.40 Design Choices and Limitations"),
            (
                "N",
                "Object detection was preferred over whole-image classification because bounding boxes and ROI checks need location information. "
                "COCO-SSD was preferred for the internship window because it runs in TensorFlow.js and already includes the dog class.",
            ),
            (
                "N",
                "Known limitations are lighting sensitivity, weaker performance for very distant dogs, and dependence on network quality for "
                "smooth streaming. Future work may use a custom local dataset, improved night support, event logs, and weatherproof housing.",
            ),
            (
                "N",
                "Even with these limits, the prototype meets the training goal: a demonstrable IoT and computer-vision alert pipeline for "
                "stray-dog monitoring near parked vehicles, built on Machine Learning concepts studied during the internship.",
            ),
            ("H", "4.41 Report Structure Alignment"),
            (
                "N",
                "Documentation in this report covers company background, training topics, project methodology, tools, hardware and software "
                "design, implementation work, figures, testing, and conclusion so that the internship work is complete and reviewable.",
            ),
            (
                "N",
                "Care was taken to keep the report structure aligned with the institute format: declaration, certificate, acknowledgement, "
                "abstract, contents, list of figures, list of tables, chapters, and references.",
            ),
            (
                "N",
                "The figures included for the project are conceptual diagrams of ESP32-CAM hardware, system architecture, detection workflow, "
                "and UI/ROI overlay. They support understanding of the implemented design without changing the original report formatting style.",
            ),
        ]
        for kind, text in blocks:
            if kind == "H":
                cur = add_heading_after(cur, text)
            else:
                cur = add_body_after(cur, text)

        # If still short, add one more moderate section (not endless notes)
        while sum(len((p.text or "").split()) for p in doc.paragraphs) < target:
            extra = [
                "Additional validation was performed by checking that person or car detections do not activate the relay path, while dog "
                "detections inside the selected zone do activate it according to the hold timer.",
                "Power stability of ESP32-CAM was treated as a practical requirement because weak supply can cause camera init failure or reboot. "
                "A stable 5V source and correct upload wiring were part of the setup checklist.",
                "The hybrid design—browser inference plus ESP32 actuation—kept the microcontroller free for streaming and GPIO work, which "
                "improved reliability compared with attempting heavy neural inference on the board itself during this training scope.",
                "Confidence threshold near 55 percent was selected after trial observations as a balance between missed dogs and false alerts. "
                "Operators can still visually confirm boxes and scores on the monitor before trusting long unattended runs.",
                "In residential use-case terms, the system is intended as an assistive monitor and alert signal generator. Any deterrent load "
                "on the relay must follow electrical safety and local rules; the internship focus remained on detection correctness and control signaling.",
                "Mapping back to training lectures, classification confidence, supervised learning intuition, and deep-learning based perception "
                "were all visible in the way COCO-SSD outputs were filtered and converted into actions.",
                "Thus the project chapter records both what was built and why the chosen architecture was suitable for a short industrial-training timeline.",
            ]
            # append each once
            for t in extra:
                cur = add_body_after(cur, t)
            break

    doc.save(str(SRC))

    doc = Document(str(SRC))
    print("VERIFY:")
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if i > 100:
            break
        if any(
            t.startswith(x)
            for x in (
                "ABSTRACT",
                "CONTENTS",
                "Candidate",
                "Certificate",
                "Acknowledgement",
                "Abstract",
                "Contents",
                "List of",
                "Chapter",
                "References",
                "CHAPTER 1",
            )
        ):
            st = p.style.name if p.style else ""
            print(f"{i:3d} [{st}] {t[:95]}")
    words = sum(len((p.text or "").split()) for p in doc.paragraphs)
    print("words", words, "pages@240", round(words / 240, 1))
    # confirm tables still after LoF
    body = doc.element.body
    print("tables", len(doc.tables))


if __name__ == "__main__":
    main()
