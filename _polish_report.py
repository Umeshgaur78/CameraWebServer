"""Final polish: line spacing, page-break hints, TOC/LoF page numbers."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = Path(r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx")
BODY_PT = 14


def set_run_font(run, size_pt=BODY_PT, bold=None):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
        set_run_font(paragraph.runs[0], BODY_PT)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, BODY_PT)


def ensure_page_break_before(paragraph):
    """Add page break at start of paragraph if not already present."""
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    # check existing pageBreakBefore
    for child in pPr:
        if child.tag == qn("w:pageBreakBefore"):
            return
    pb = OxmlElement("w:pageBreakBefore")
    pPr.append(pb)


def main():
    doc = Document(str(SRC))

    # Apply 1.5 line spacing to Normal body paragraphs (keeps font size)
    try:
        normal = doc.styles["Normal"]
        if normal.font.size is None or normal.font.size.pt == 0:
            pass
        # don't change font name/size/color; only spacing if empty-ish
        pf = normal.paragraph_format
        if pf.line_spacing is None:
            pf.line_spacing = 1.5
    except Exception:
        pass

    for p in doc.paragraphs:
        st = p.style.name if p.style else ""
        t = (p.text or "").strip()
        if st == "Normal" and t:
            try:
                if p.paragraph_format.line_spacing is None:
                    p.paragraph_format.line_spacing = 1.5
            except Exception:
                pass

    # Page break before major headings (common report format)
    markers = (
        "CERTIFICATE",
        "ACKNOWLEDGEMENT",
        "ABSTRACT",
        "List of Figures",
        "List of Tables",
        "CHAPTER 1:",
        "CHAPTER 2:",
        "CHAPTER 3:",
        "CHAPTER 4:",
        "CHAPTER 5:",
        "REFERENCES",
    )
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        for m in markers:
            if t.startswith(m):
                ensure_page_break_before(p)
                break

    # TOC page number refresh for ~30 pages with figures/spacing
    toc = {
        "Chapter 1: Introduction": "Chapter 1: Introduction ............................................................................................................. 1",
        "Chapter 2: Company Infrastructure": "Chapter 2: Company Infrastructure ............................................................................................. 4",
        "Chapter 3: Training Attended": "Chapter 3: Training Attended ...................................................................................................... 7",
        "Chapter 4: Systems / Project Development": "Chapter 4: Systems / Project Development ................................................................................. 12",
        "Chapter 5: Results and Conclusion": "Chapter 5: Results and Conclusion ............................................................................................. 27",
        "References": "References ................................................................................................................................. 30",
        "Candidate's Declaration": "Candidate's Declaration ......................................................................................................... i",
        "Certificate": "Certificate .............................................................................................................................. ii",
        "Acknowledgement": "Acknowledgement ................................................................................................................ iii",
        "Abstract": "Abstract ................................................................................................................................. iv",
        "List of Figures": "List of Figures ........................................................................................................................ v",
        "List of Tables": "List of Tables ........................................................................................................................ vi",
    }
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        for key, val in toc.items():
            if t.startswith(key) and ("..." in t or t.startswith("Chapter") or t.startswith("References") or t.startswith("Candidate") or t.startswith("Certificate") or t.startswith("Acknowledgement") or t.startswith("Abstract") or t.startswith("List of")):
                # only update dotted TOC-like lines or chapter toc lines
                if "..." in t or t.startswith("Chapter") or t.startswith("References"):
                    replace_paragraph_text(p, val)
                break

    if len(doc.tables) >= 1:
        t0 = doc.tables[0]
        while len(t0.rows) < 5:
            t0.add_row()
        data = [
            ("Fig. 4.1", "ESP32-CAM hardware concept", "18"),
            ("Fig. 4.2", "System architecture of dog detection alert system", "19"),
            ("Fig. 4.3", "Detection and relay-hold workflow", "20"),
            ("Fig. 4.4", "Monitoring UI with DETECT ZONE and dog box", "21"),
        ]
        for i, (a, b, c) in enumerate(data, start=1):
            t0.rows[i].cells[0].text = a
            t0.rows[i].cells[1].text = b
            t0.rows[i].cells[2].text = c

    if len(doc.tables) >= 2:
        t1 = doc.tables[1]
        while len(t1.rows) < 3:
            t1.add_row()
        t1.rows[1].cells[0].text = "Table 4.1"
        t1.rows[1].cells[1].text = "Tools and technologies used in the project"
        t1.rows[1].cells[2].text = "17"
        if len(t1.rows) > 2:
            t1.rows[2].cells[0].text = "Table 4.2"
            t1.rows[2].cells[1].text = "GPIO pin mapping (Relay/Flash)"
            t1.rows[2].cells[2].text = "22"

    doc.save(str(SRC))

    words = sum(len((p.text or "").split()) for p in doc.paragraphs)
    print("done words", words, "pages@250w", round(words / 250, 1))
    # abstract snippet
    for p in doc.paragraphs:
        if "complete real-time dog detection" in (p.text or ""):
            print("abstract_updated_ok")
            break


if __name__ == "__main__":
    main()
