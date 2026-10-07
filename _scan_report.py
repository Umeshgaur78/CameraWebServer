from docx import Document

p = r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx"
d = Document(p)
print("=== ABSTRACT AREA ===")
for i, para in enumerate(d.paragraphs):
    t = (para.text or "").strip()
    if "ABSTRACT" in t.upper() or "abstract" in t.lower() or "progress" in t.lower() or "under development" in t.lower() or "successfully" in t.lower():
        print(i, repr(t[:120]))
print("=== CHAPTER 4/5 HEADERS ===")
for i, para in enumerate(d.paragraphs):
    t = (para.text or "").strip()
    if t.startswith("Chapter") or t.startswith("4.") or t.startswith("5.") or "LIST OF" in t.upper() or t.startswith("CONTENTS"):
        print(i, repr(t[:100]))
