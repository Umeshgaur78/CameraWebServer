from docx import Document
import os

path = r"c:\Users\umesh\OneDrive\Desktop\report vinay sir.docx"
d = Document(path)

print("=== ALL NON-EMPTY PARAS ===")
for i, p in enumerate(d.paragraphs):
    t = (p.text or "").strip()
    has_img = bool(p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
    if not t:
        if has_img:
            print(i, "[IMAGE]")
        continue
    style = p.style.name if p.style else ""
    print(f"{i}|{style}|{t[:160]}")

print("\n=== TABLES ===")
for ti, t in enumerate(d.tables):
    print("table", ti, "rows", len(t.rows), "cols", len(t.columns))
    for r in t.rows[:8]:
        print(" | ".join(c.text.strip()[:50] for c in r.cells))

print("\n=== AROUND CERTIFICATE ===")
for i in range(39, min(70, len(d.paragraphs))):
    p = d.paragraphs[i]
    t = (p.text or "").strip()
    has_img = bool(p._p.xpath('.//*[local-name()="drawing" or local-name()="pict"]'))
    print(i, repr(t[:80]), "img" if has_img else "")

print("\nfile size", os.path.getsize(path))
print("total paras", len(d.paragraphs))
